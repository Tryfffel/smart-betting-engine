"""Skriver leads_master.xlsx + daglig_logg.csv + per-lead Markdown."""
from __future__ import annotations

import csv
import logging
import re
from datetime import date
from pathlib import Path
from string import Template
from typing import Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models import Lead

logger = logging.getLogger(__name__)

OUTPUTS = Path(__file__).resolve().parent.parent / "outputs"

COLUMNS = [
    "score",
    "organisation",
    "org_typ",
    "roll",
    "publicerad",
    "sista_ansokningsdag",
    "ansvarig_rekryterare",
    "kontakt_namn",
    "kontakt_titel",
    "kontakt_email",
    "kontakt_kalla",
    "annons_url",
    "annons_id",
    "platsbanken",
    "linkedin_finns",
    "motivering",
    "status",
]

SCORE_ORDER = {"A": 0, "B": 1, "C": 2}

MD_TEMPLATE = Template(
    """# $roll – $organisation
**Score:** $score | **Publicerad:** $publicerad | **Sista ansökan:** $sista_ansokan

## Annonsen
- Länk: $annons_url
- Typ: $org_typ
- Källa: $source
- Beskrivning: $beskrivning

## Kontaktperson
- Namn: $kontakt_namn
- Titel: $kontakt_titel
- E-post: $kontakt_email
- Källa: $kontakt_kalla

## Motivering till score
$motivering

## Pitchnotering
$pitch
"""
)


def _lead_row(lead: Lead) -> list:
    platsbanken = lead.annons_url if "arbetsformedlingen" in lead.annons_url else ""
    return [
        lead.score,
        lead.organisation,
        lead.org_typ,
        lead.roll,
        lead.publicerad.isoformat(),
        lead.sista_ansokningsdag.isoformat() if lead.sista_ansokningsdag else "",
        lead.ansvarig_rekryterare or "",
        lead.kontakt.namn or "",
        lead.kontakt.titel or "",
        lead.kontakt.email or "",
        lead.kontakt.kalla or "",
        lead.annons_url,
        lead.annons_id,
        platsbanken,
        "ja" if lead.linkedin_finns else "",
        lead.motivering,
        lead.status,
    ]


def _sorted_leads(leads: Iterable[Lead]) -> list[Lead]:
    return sorted(leads, key=lambda l: (SCORE_ORDER.get(l.score, 9), -_pub_ord(l.publicerad)))


def _pub_ord(d: date) -> int:
    return d.toordinal()


def _slugify(text: str) -> str:
    s = re.sub(r"[^\w\s-]", "", text, flags=re.UNICODE).strip().lower()
    return re.sub(r"[-\s]+", "-", s)[:80]


def _pitch_for(lead: Lead) -> str:
    """Enkel regelbaserad pitch-notering — finjusteras senare."""
    parts = []
    if lead.org_typ == "kommun":
        parts.append(
            "Kommunchefer och förvaltningsledning är överrepresenterade bland "
            "Dagens Samhälles 125 000 läsare/vecka."
        )
    elif lead.org_typ == "region":
        parts.append(
            "Beslutsfattare i regioner läser Dagens Samhälle veckovis — "
            "Digital + Sociala medier-paketet (~22 400 kr) ger bred täckning."
        )
    elif lead.org_typ == "myndighet":
        parts.append(
            "Myndighetschefer och stabsfunktioner nås effektivt via Native (29 900) "
            "eller Banner (14 900)."
        )
    elif lead.org_typ == "kommunalt_bolag":
        parts.append(
            "Kommunala bolag har ofta utrymme i kommunikationsbudget — "
            "Print + Digital-kombo (~25 000) brukar bita."
        )
    elif lead.org_typ == "rekryteringsbyra":
        parts.append(
            "Rekryteringsbyrå för offentlig sektor — direkt köpare. Hänvisa till "
            "kommun-/region-paket. Kontorets seniora konsult är ofta beslutsfattare."
        )
    if lead.sista_ansokningsdag:
        parts.append(f"Sista ansökan: {lead.sista_ansokningsdag.isoformat()} — agera snabbt.")
    return " ".join(parts) or "Ingen automatisk pitch."


def write_outputs(leads: list[Lead], run_date: date) -> Path:
    """Skriver alla outputs för en körning. Returnerar mapp-stigen."""
    sorted_leads = _sorted_leads(leads)
    run_dir = OUTPUTS / run_date.isoformat()
    run_dir.mkdir(parents=True, exist_ok=True)
    (run_dir / "a_leads").mkdir(exist_ok=True)

    _write_xlsx(sorted_leads, run_dir / "leads_master.xlsx")
    _write_csv(sorted_leads, run_dir / "daglig_logg.csv")
    _write_a_lead_markdown(sorted_leads, run_dir / "a_leads")

    logger.info("Export: %d leads skrivna till %s", len(sorted_leads), run_dir)
    return run_dir


def _write_xlsx(leads: list[Lead], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Leads"

    # Header
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3B3B3B")
    for col, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")

    for row_idx, lead in enumerate(leads, start=2):
        for col_idx, value in enumerate(_lead_row(lead), start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Conditional formatting på score-kolumnen (kolumn A)
    last_row = max(2, len(leads) + 1)
    score_range = f"A2:A{last_row}"
    ws.conditional_formatting.add(
        score_range,
        CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", fgColor="C6EFCE")),
    )
    ws.conditional_formatting.add(
        score_range,
        CellIsRule(operator="equal", formula=['"B"'], fill=PatternFill("solid", fgColor="FFEB9C")),
    )
    ws.conditional_formatting.add(
        score_range,
        CellIsRule(operator="equal", formula=['"C"'], fill=PatternFill("solid", fgColor="D9D9D9")),
    )

    # Kolumnbredd, frys + autofilter
    widths = {
        1: 8, 2: 28, 3: 16, 4: 28, 5: 12, 6: 12, 7: 22, 8: 22,
        9: 22, 10: 28, 11: 28, 12: 48, 13: 14, 14: 28, 15: 12, 16: 60, 17: 10,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    wb.save(path)


def _write_csv(leads: list[Lead], path: Path) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(COLUMNS)
        for lead in leads:
            writer.writerow(_lead_row(lead))


def _write_a_lead_markdown(leads: list[Lead], dir_path: Path) -> None:
    for lead in leads:
        if lead.score != "A":
            continue
        slug = f"{_slugify(lead.organisation)}_{_slugify(lead.roll)}"
        path = dir_path / f"{slug}.md"
        content = MD_TEMPLATE.substitute(
            roll=lead.roll,
            organisation=lead.organisation,
            score=lead.score,
            publicerad=lead.publicerad.isoformat(),
            sista_ansokan=lead.sista_ansokningsdag.isoformat() if lead.sista_ansokningsdag else "okänt",
            annons_url=lead.annons_url,
            org_typ=lead.org_typ,
            source=lead.source,
            beskrivning=(lead.beskrivning or "")[:600].replace("\n", " "),
            kontakt_namn=lead.kontakt.namn or "—",
            kontakt_titel=lead.kontakt.titel or "—",
            kontakt_email=lead.kontakt.email or "—",
            kontakt_kalla=lead.kontakt.kalla or "ej hittad",
            motivering=lead.motivering,
            pitch=_pitch_for(lead),
        )
        path.write_text(content, encoding="utf-8")
